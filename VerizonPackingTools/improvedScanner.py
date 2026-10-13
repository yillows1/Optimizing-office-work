import win32com.client
import os
import io
import json
import tempfile
import re
import difflib
import sys
from collections import defaultdict
from datetime import datetime, timedelta

stdout_reconfigure = getattr(sys.stdout, "reconfigure", None)
if callable(stdout_reconfigure):
    stdout_reconfigure(errors="replace")

# --- Config ---
LIST_FILE = "receivedList.txt"

# Fuzzy matching settings
FUZZY_ENABLED = True
FUZZY_THRESHOLD = 0.85
FUZZY_MIN_LEN = 6

RITM_RE = re.compile(r"\bRITM[\s:#-]*\d+\b", re.IGNORECASE)
WORD_RE = re.compile(r"[A-Za-z0-9]+")


def _range_from_choice(choice: str):
    now = datetime.now()
    if choice == "1":
        return now - timedelta(days=90), "Past 3 months"
    if choice == "3":
        return now - timedelta(days=365), "Past 12 months"
    if choice == "4":
        return None, "All time"
    return now - timedelta(days=182), "Past 6 months"


def read_attachment_bytes(attachment):
    try:
        suffix = os.path.splitext(attachment.FileName)[1]
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
            tmp_path = tmp.name
        attachment.SaveAsFile(tmp_path)
        with open(tmp_path, "rb") as f:
            data = f.read()
        os.remove(tmp_path)
        return data
    except Exception as e:
        print(f"  [!] Could not read {attachment.FileName}: {e}")
        return None


def extract_text_from_bytes(data, filename):
    ext = os.path.splitext(filename)[1].lower()
    try:
        if ext in (".txt", ".csv", ".log", ".json", ".xml", ".html", ".md"):
            return data.decode("utf-8", errors="ignore")
        elif ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        elif ext == ".docx":
            from docx import Document
            doc = Document(io.BytesIO(data))
            return "\n".join(p.text for p in doc.paragraphs)
        elif ext == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), data_only=True)
            parts = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    parts.append(" ".join(str(c) for c in row if c is not None))
            return "\n".join(parts)
        else:
            return data.decode("latin-1", errors="ignore")
    except Exception as e:
        print(f"  [!] Could not parse {filename}: {e}")
        return ""


def find_terms_in_text(text, search_terms_lower):
    if not text:
        return set()
    lowered = text.lower()
    hits = set()
    for term, term_lower in search_terms_lower:
        if term_lower in lowered:
            hits.add(term)
    return hits


def find_ritms_in_text(text):
    return sorted({
        re.sub(r"[^A-Za-z0-9]", "", match).upper()
        for match in RITM_RE.findall(text or "")
    })


def _looks_like_address_line(line: str) -> bool:
    """Return True if a line looks like part of an address.

    Address material almost always has a comma or a digit, or is a short
    run of capitalized tokens (e.g. 'MA', 'USA', 'Suite 200'). Normal prose
    like 'Please communicate with the below mentioned Contact Person...'
    has none of these, so it gets rejected.
    """
    stripped = line.strip()
    if not stripped:
        return False
    if "," in stripped or re.search(r"\d", stripped):
        return True
    words = stripped.split()
    if len(words) <= 3 and all(w[:1].isupper() or w.isupper() for w in words):
        return True
    return False


def find_location_address_in_text(text):
    lines = (text or "").splitlines()
    for index, line in enumerate(lines):
        match = re.search(r"\bLocation\b\s*:?\s*(.*)", line, re.IGNORECASE)
        if not match:
            continue

        address = match.group(1).strip()
        for following_line in lines[index + 1:index + 4]:
            # Stop as soon as we have a street and a city.
            address_parts = [part.strip() for part in address.split(",")]
            if len(address_parts) >= 2 and address_parts[0] and address_parts[1]:
                break
            following_line = following_line.strip()
            if not following_line:
                continue
            if re.search(r"\b[A-Za-z][A-Za-z ]*\s*:", following_line):
                break
            address = f"{address} {following_line}".strip()

        parts = [part.strip() for part in address.split(",")]
        if len(parts) < 2 or not parts[0] or not parts[1]:
            continue
        return ", ".join(parts)
    return ""


def find_location_in_text(text):
    parts = [part.strip() for part in find_location_address_in_text(text).split(",")]
    if len(parts) < 2 or not parts[0] or not parts[1]:
        return ""
    city = parts[1]
    if len(parts) >= 3 and parts[2]:
        state_match = re.match(r"([A-Za-z]{2})\b", parts[2])
        if state_match:
            return f"{city} {state_match.group(1).upper()}"
    return city

def _build_words_by_len(text):
    words = set(WORD_RE.findall(text))
    words_by_len = defaultdict(list)
    for w in words:
        w_low = w.lower()
        words_by_len[len(w_low)].append((w, w_low))
    return words_by_len


def find_near_misses_from_words_by_len(words_by_len, unmatched_terms):
    result = {}
    if not words_by_len or not unmatched_terms:
        return result
    matcher = difflib.SequenceMatcher()
    for term in unmatched_terms:
        if len(term) < FUZZY_MIN_LEN:
            continue
        term_lower = term.lower()
        term_len = len(term_lower)
        candidates = []
        for delta in (-2, -1, 0, 1, 2):
            candidates.extend(words_by_len.get(term_len + delta, ()))
        if not candidates:
            continue
        matcher.set_seq2(term_lower)
        hits = []
        for w_orig, w_low in candidates:
            matcher.set_seq1(w_low)
            ratio = matcher.ratio()
            if ratio >= FUZZY_THRESHOLD:
                hits.append((w_orig, ratio))
        if hits:
            hits.sort(key=lambda x: -x[1])
            result[term] = hits[:3]
    return result


def get_folders(folder):
    yield folder
    for sub in folder.Folders:
        yield from get_folders(sub)


def run_scan(time_range_choice, search_terms, log_path, json_results_path=None):
    search_terms = [t.strip() for t in search_terms if t and t.strip()]
    if not search_terms:
        raise ValueError("No search terms provided.")
    search_terms_lower = [(t, t.lower()) for t in search_terms]
    cutoff, range_label = _range_from_choice(time_range_choice)

    print(f"\nTime range:   {range_label}")
    if cutoff:
        print(f"Cutoff date:  {cutoff.strftime('%Y-%m-%d')}")
    print(f"Search terms: {len(search_terms)}")
    for t in search_terms:
        print(f"  - {t}")
    print()

    outlook = win32com.client.Dispatch("Outlook.Application").GetNamespace("MAPI")
    inbox = outlook.GetDefaultFolder(6)

    matches_by_email = defaultdict(lambda: {"terms": set(), "files": set()})
    fuzzy_by_email = defaultdict(lambda: defaultdict(set))
    found_terms = set()
    scanned_count = 0
    skipped_count = 0
    gui_exact_matches = []
    gui_near_matches = []

    folders_to_scan = list(get_folders(inbox))
    print(f"Scanning Inbox + all subfolders ({len(folders_to_scan)} folder(s))...\n")

    for folder in folders_to_scan:
        messages = folder.Items
        try:
            messages.Sort("[ReceivedTime]", True)
        except Exception:
            pass
        if messages.Count == 0:
            continue

        folder_scanned = 0
        folder_skipped = 0
        hit_cutoff = False

        for message in messages:
            if cutoff is not None:
                try:
                    received = message.ReceivedTime
                    if received.year < 1900:
                        skipped_count += 1
                        folder_skipped += 1
                        continue
                    received_dt = datetime(
                        received.year, received.month, received.day,
                        received.hour, received.minute, received.second,
                    )
                    if received_dt < cutoff:
                        hit_cutoff = True
                        break
                except Exception:
                    skipped_count += 1
                    folder_skipped += 1
                    continue

            try:
                if message.Attachments.Count == 0:
                    continue
            except Exception:
                continue

            scanned_count += 1
            folder_scanned += 1
            email_terms = set()
            email_files = set()
            need_json = bool(json_results_path)

            if need_json:
                try:
                    email_subject = str(message.Subject or "")
                except Exception:
                    email_subject = ""
                try:
                    email_body = str(message.Body or "")
                except Exception:
                    email_body = ""
                try:
                    sender_name = str(message.SenderName or "")
                except Exception:
                    sender_name = ""
                try:
                    sender_email = str(message.SenderEmailAddress or "")
                except Exception:
                    sender_email = ""
                try:
                    entry_id = str(message.EntryID or "")
                except Exception:
                    entry_id = ""
                try:
                    store_id = str(folder.StoreID or "")
                except Exception:
                    store_id = ""
                received_text = str(message.ReceivedTime)
                subject_ritms = find_ritms_in_text(email_subject)
                email_ritms = find_ritms_in_text(f"{email_subject}\n{email_body}")

            attachment_records = []
            for attachment in message.Attachments:
                filename = attachment.FileName
                data = read_attachment_bytes(attachment)
                if data is None:
                    continue
                content = extract_text_from_bytes(data, filename)
                hits = find_terms_in_text(content, search_terms_lower)
                attachment_records.append({
                    "filename": filename,
                    "content": content,
                    "hits": hits,
                    "content_len": len(content),
                })
                if hits:
                    email_terms |= hits
                    email_files.add(filename)
                    print(f"  {filename} -> {', '.join(sorted(hits))}")

            if FUZZY_ENABLED and attachment_records:
                try:
                    entry_key = message.EntryID
                except Exception:
                    entry_key = f"{message.Subject}|{message.ReceivedTime}"
                for rec in attachment_records:
                    content = rec["content"]
                    if not content:
                        continue
                    exact_here = {t.casefold() for t in rec["hits"]}
                    unmatched_for_file = [
                        t for t in search_terms if t.casefold() not in exact_here
                    ]
                    if not unmatched_for_file:
                        continue
                    words_by_len = _build_words_by_len(content)
                    near = find_near_misses_from_words_by_len(
                        words_by_len, unmatched_for_file
                    )
                    if not near:
                        continue
                    for term, guess_hits in near.items():
                        for guess_word, ratio in guess_hits:
                            fuzzy_by_email[entry_key][rec["filename"]].add(
                                (term, guess_word, round(ratio, 3))
                            )
                    if need_json:
                        for term, guess_hits in near.items():
                            for guess_word, ratio in guess_hits:
                                gui_near_matches.append({
                                    "term": term,
                                    "guess": guess_word,
                                    "similarity": round(ratio, 3),
                                    "filename": rec["filename"],
                                    "subject": email_subject,
                                    "folder": folder.Name,
                                    "received": received_text,
                                    "subject_ritms": subject_ritms,
                                })

            if need_json:
                for rec in attachment_records:
                    if not rec["hits"]:
                        continue
                    content = rec["content"]
                    attachment_ritms = find_ritms_in_text(content)
                    location = find_location_in_text(content)
                    location_address = find_location_address_in_text(content)
                    for term in rec["hits"]:
                        gui_exact_matches.append({
                            "term": term,
                            "filename": rec["filename"],
                            "subject": email_subject,
                            "sender": sender_name,
                            "sender_email": sender_email,
                            "entry_id": entry_id,
                            "store_id": store_id,
                            "folder": folder.Name,
                            "received": received_text,
                            "ritms": sorted(set(attachment_ritms) | set(email_ritms)),
                            "subject_ritms": subject_ritms,
                            "location": location,
                            "location_address": location_address,
                        })

            if email_terms:
                try:
                    key = message.EntryID
                except Exception:
                    key = f"{message.Subject}|{message.ReceivedTime}"
                entry = matches_by_email[key]
                entry["subject"] = message.Subject
                entry["sender"] = message.SenderName
                entry["received"] = message.ReceivedTime
                entry["folder"] = folder.Name
                entry["terms"] |= email_terms
                entry["files"] |= email_files
                found_terms |= email_terms

            if folder_scanned % 50 == 0:
                print(f"  ...scanned {folder_scanned} email(s) in [{folder.Name}]")

        if folder_scanned or folder_skipped:
            print(
                f"[{folder.Name}] scanned {folder_scanned}, skipped {folder_skipped}"
                + (" (hit cutoff, stopped early)" if hit_cutoff else "")
            )

    multi_term_emails = [e for e in matches_by_email.values() if len(e["terms"]) >= 2]
    single_term_emails = [e for e in matches_by_email.values() if len(e["terms"]) == 1]

    fuzzy_terms_anywhere = set()
    for fmap in fuzzy_by_email.values():
        for hint_set in fmap.values():
            for term, _, _ in hint_set:
                fuzzy_terms_anywhere.add(term)

    missing_terms = [t for t in search_terms if t not in found_terms]

    lines = []
    def w(s=""):
        lines.append(s)
        print(s)

    w()
    w("=" * 60)
    w(f"SEARCH SUMMARY - {len(search_terms)} term(s) searched")
    w("=" * 60)
    w(f"Time range:                  {range_label}")
    if cutoff:
        w(f"Cutoff date:                 {cutoff.strftime('%Y-%m-%d')}")
    w(f"Emails scanned (in range):   {scanned_count}")
    w(f"Emails skipped:              {skipped_count}")
    w(f"Emails with 2+ terms:        {len(multi_term_emails)}")
    w(f"Emails with exactly 1 term:  {len(single_term_emails)}")
    w(f"Terms never found:           {len(missing_terms)}")
    w(f"Fuzzy matching:              {'ON' if FUZZY_ENABLED else 'OFF'}")
    w()

    w("=" * 60)
    w("SECTION 1: MULTIPLE TERMS IN THE SAME EMAIL")
    w("=" * 60)
    if not multi_term_emails:
        w("(none)")
    else:
        for e in sorted(multi_term_emails, key=lambda x: x["received"], reverse=True):
            w(f"\n {e['subject']}")
            w(f"   From:     {e['sender']}")
            w(f"   Received: {e['received']}")
            w(f"   Folder:   {e['folder']}")
            w(f"   Terms:    {', '.join(sorted(e['terms']))}")
            w(f"   Files:    {', '.join(sorted(e['files']))}")

    w()
    w("=" * 60)
    w("SECTION 2: SINGLE TERM MATCHES")
    w("=" * 60)
    if not single_term_emails:
        w("(none)")
    else:
        for e in sorted(single_term_emails, key=lambda x: x["received"], reverse=True):
            w(f"\n {e['subject']}")
            w(f"   From:     {e['sender']}")
            w(f"   Received: {e['received']}")
            w(f"   Folder:   {e['folder']}")
            w(f"   Term:     {', '.join(sorted(e['terms']))}")
            w(f"   Files:    {', '.join(sorted(e['files']))}")

    w()
    w("=" * 60)
    w("SECTION 3: TERMS NEVER FOUND")
    w("=" * 60)
    if not missing_terms:
        w("(none - every term was found at least once)")
    else:
        for t in missing_terms:
            w(f"  X {t}")

    w()
    w("=" * 60)
    w("SECTION 4: GROUPED BY CO-OCCURRENCE")
    w("=" * 60)
    w()

    groups = defaultdict(list)
    for entry_id, e in matches_by_email.items():
        groups[frozenset(e["terms"])].append((entry_id, e))

    def group_sort_key(item):
        terms, emails = item
        newest = max(e["received"] for _, e in emails)
        return (-len(terms), newest)

    ordered_groups = sorted(groups.items(), key=group_sort_key)

    def files_for_group(emails):
        files = set()
        for _, e in emails:
            files |= e["files"]
        return sorted(files)

    letter = ord("A")
    for terms, emails in ordered_groups:
        files = files_for_group(emails)
        file_label = ", ".join(files) if files else "(unknown attachment)"
        w(f"{chr(letter)}) {file_label}")
        for t in sorted(terms):
            w(f"   {t}")

        terms_lower = {t.lower() for t in terms}
        suggestion_map = defaultdict(set)
        for entry_id, _ in emails:
            fmap = fuzzy_by_email.get(entry_id)
            if not fmap:
                continue
            for filename, hint_set in fmap.items():
                if filename in files:
                    for term, guess, ratio in hint_set:
                        if term.lower() in terms_lower:
                            continue
                        if guess.lower() in terms_lower:
                            continue
                        suggestion_map[filename].add((term, guess, ratio))

        if suggestion_map:
            w()
            w("   ! Possible near-miss (needs human check):")
            for filename in sorted(suggestion_map):
                for term, guess, ratio in sorted(suggestion_map[filename]):
                    w(f"     {filename}: '{guess}' ~ {term}  (similarity {ratio})")
        w()
        letter += 1

    truly_missing = [
        t for t in search_terms
        if t not in found_terms and t not in fuzzy_terms_anywhere
    ]
    if truly_missing:
        w(f"{chr(letter)}) N/A")
        for t in sorted(truly_missing):
            w(f"   {t}")
        w()
        letter += 1

    with open(log_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))

    if json_results_path:
        exact_ritms = set()
        for entry in gui_exact_matches:
            for ritm in entry.get("subject_ritms") or []:
                exact_ritms.add(ritm.casefold())
            for ritm in entry.get("ritms") or []:
                exact_ritms.add(ritm.casefold())

        if exact_ritms:
            anchored = [
                entry for entry in gui_near_matches
                if {r.casefold() for r in (entry.get("subject_ritms") or [])} & exact_ritms
            ]
        else:
            anchored = []

        deduped = {}
        for entry in anchored:
            key = (
                entry["term"].casefold(),
                entry["folder"].casefold(),
                entry["filename"].casefold(),
            )
            existing = deduped.get(key)
            if existing is None or entry["similarity"] > existing["similarity"]:
                deduped[key] = entry

        near_matches_out = sorted(
            deduped.values(),
            key=lambda e: (
                e["received"],
                e["folder"].casefold(),
                e["filename"].casefold(),
                e["term"].casefold(),
                -e["similarity"],
            ),
        )

        with open(json_results_path, "w", encoding="utf-8") as f:
            json.dump(
                {"exact_matches": gui_exact_matches, "near_matches": near_matches_out},
                f, ensure_ascii=False, indent=2,
            )

    print(f"\nReport saved to: {log_path}")


def _cli_main():
    log_file = os.environ.get("SCAN_RESULTS_OUTPUT", "scan_results.txt")
    json_file = os.environ.get("SCAN_RESULTS_JSON")
    search_terms = [
        term.strip()
        for term in os.environ.get("SCAN_SEARCH_TERMS", "").splitlines()
        if term.strip()
    ]
    if not search_terms:
        if not os.path.exists(LIST_FILE):
            print(f"Could not find {LIST_FILE} in the current folder.")
            input("Press Enter to exit...")
            sys.exit(1)
        with open(LIST_FILE, "r", encoding="utf-8") as f:
            search_terms = [line.strip() for line in f if line.strip()]
    if not search_terms:
        print(f"{LIST_FILE} is empty.")
        input("Press Enter to exit...")
        sys.exit(1)

    print("\nChoose a time range to scan:")
    print("  (1) Past 3 months")
    print("  (2) Past 6 months")
    print("  (3) Past 12 months")
    print("  (4) All time")
    choice = input("Selection [2]: ").strip() or "2"
    try:
        run_scan(choice, search_terms, log_file, json_file)
    except Exception as e:
        print(f"Scan failed: {e}")
        input("\nPress Enter to exit...")
        sys.exit(1)
    input("\nPress Enter to exit...")


if __name__ == "__main__":
    _cli_main()